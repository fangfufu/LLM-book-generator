import pathlib
import tempfile
import unittest

from docx import Document

from src.emitter import html_to_docx, markdown_to_docx
from src.processor import (
    clean_html_tree,
    clean_markdown_text,
    convert_markdown_to_raw_html,
    process_markdown_to_html,
)


class TestProcessor(unittest.TestCase):
    def test_clean_markdown_text(self):
        raw = "Heading\n\n\n\n\nParagraph text.\n\n\n\n"
        cleaned = clean_markdown_text(raw)
        self.assertEqual(cleaned, "Heading\n\nParagraph text.")

    def test_convert_markdown_to_raw_html(self):
        markdown_text = "# Chapter 1\n\nThis is a **bold** and *italic* paragraph."
        _, html_content = convert_markdown_to_raw_html(markdown_text)
        self.assertIn("<h1>Chapter 1</h1>", html_content)
        self.assertIn("<strong>bold</strong>", html_content)
        self.assertIn("<em>italic</em>", html_content)

    def test_clean_html_tree_removes_empty_and_br_only_paragraphs(self):
        html_input = (
            "<p>Valid text</p>"
            "<p></p>"
            "<p><br></p>"
            "<p><br><br/></p>"
            "<p>Another paragraph</p>"
        )
        cleaned_tree = clean_html_tree(html_input)
        p_tags = cleaned_tree.xpath(".//p")
        self.assertEqual(len(p_tags), 2)
        self.assertEqual(p_tags[0].text_content(), "Valid text")
        self.assertEqual(p_tags[1].text_content(), "Another paragraph")

    def test_save_intermediate_html_when_enabled(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            config = {
                "cache_dir": tmp_dir,
                "debug_options": {"save_intermediate_html": True},
            }
            process_markdown_to_html(
                "## Test section\n\nHello world.",
                config=config,
                context_label="test_ctx",
            )
            debug_dir = pathlib.Path(tmp_dir) / "debug_html"
            self.assertTrue(debug_dir.exists())
            html_files = list(debug_dir.glob("*.html"))
            self.assertEqual(len(html_files), 1)
            self.assertIn("Hello world.", html_files[0].read_text(encoding="utf-8"))


class TestEmitter(unittest.TestCase):
    def test_markdown_to_docx_populates_document(self):
        doc = Document()
        config = {
            "style_params": {
                "font_name": "Calibri",
                "font_size": 11.0,
            },
            "debug_options": {},
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            markdown_to_docx(
                markdown_text="## Section One\n\nParagraph with **bold** text.",
                container_obj=doc,
                doc=doc,
                config=config,
                usable_width_inches=5.0,
                equation_image_dir=pathlib.Path(tmp_dir),
            )
            # Verify document has paragraphs
            para_texts = [p.text for p in doc.paragraphs if p.text]
            self.assertIn("Section One", para_texts)
            self.assertTrue(any("bold text" in t for t in para_texts))

    def test_html_to_docx_direct(self):
        doc = Document()
        config = {"style_params": {}}
        with tempfile.TemporaryDirectory() as tmp_dir:
            html_input = "<p>Direct HTML to docx test</p>"
            html_to_docx(
                html_tree_or_content=html_input,
                container_obj=doc,
                doc=doc,
                config=config,
                usable_width_inches=5.0,
                equation_image_dir=pathlib.Path(tmp_dir),
            )
            para_texts = [p.text for p in doc.paragraphs]
            self.assertIn("Direct HTML to docx test", para_texts)
